import json
import os
import tempfile
import unittest
from pathlib import Path

from wikiskill.agents import InferenceAgent
from wikiskill.benchmarks.sandbox import DockerSandbox
from wikiskill.benchmarks.spreadsheet import SpreadsheetBench
from wikiskill.data import Task
from wikiskill.demo import ScriptedModel
from wikiskill.model import ChatResponse
from wikiskill.storage import Workspace


@unittest.skipUnless(os.environ.get("WIKISKILL_DOCKER_TEST") == "1", "Set WIKISKILL_DOCKER_TEST=1 after sandbox-build")
class SandboxIntegrationTests(unittest.TestCase):
    def test_invalid_generated_workbook_is_a_failed_task(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook = openpyxl.Workbook()
            workbook.active["A1"] = 5
            workbook.save(root / "input.xlsx")
            workbook.save(root / "reference.xlsx")
            workbook.close()
            task = Task("bad-file", "Write the workbook.", "workbook", context={
                "cases": [{"input_workbook": "input.xlsx"}], "answer_position": "A1",
                "instruction_type": "Cell-Level Manipulation"}, evaluation={"reference_workbooks": ["reference.xlsx"]})
            class Model:
                identity = {"adapter": "scripted_integration"}
                step = 0
                def complete(self, messages, tools):
                    self.step += 1
                    if self.step == 1:
                        return ScriptedModel.call("bash", {"command": "printf invalid > output.xlsx"}, 1)
                    return ChatResponse({"role": "assistant", "content": "done"})
            directory = root / "execution"
            directory.mkdir()
            benchmark = SpreadsheetBench(root, {})
            result = InferenceAgent(Model(), benchmark=benchmark).run(task.public(), {}, directory=directory)
            self.assertEqual(benchmark.score("done", task, result), 0)
            self.assertEqual(result.details["cases"][0]["failure"], "invalid_output_workbook")

    def test_real_bash_formula_recalculation_scoring_and_artifact_archive(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = openpyxl.Workbook()
            source.active["B1"], source.active["C1"] = 2, 3
            source.save(root / "input.xlsx")
            source.active["A1"] = 5
            source.save(root / "reference.xlsx")
            source.close()
            directory = root / "execution"
            directory.mkdir()
            task = Task("sheet", "Sum B1 and C1 into A1.", "workbook", context={
                "cases": [{"input_workbook": "input.xlsx"}], "answer_position": "Sheet!A1",
                "instruction_type": "Cell-Level Manipulation"}, evaluation={"reference_workbooks": ["reference.xlsx"]})

            class Model:
                identity = {"adapter": "scripted_integration"}
                step = 0
                def complete(self, messages, tools):
                    self.step += 1
                    if self.step == 1:
                        command = "python - <<'PY'\nfrom pathlib import Path\nfrom openpyxl import load_workbook\nassert not Path('/workspace/reference.xlsx').exists()\nassert not Path('/workspace/task/reference.xlsx').exists()\nassert not Path('/var/run/docker.sock').exists()\nw=load_workbook('input.xlsx')\nw.active['A1']='=SUM(B1:C1)'\nw.save('output.xlsx')\nprint('saved')\nPY"
                        return ScriptedModel.call("bash", {"command": command}, 1)
                    output = json.loads(messages[-1]["content"])
                    if output["exit_code"] != 0:
                        raise RuntimeError(output["stderr"])
                    return ChatResponse({"role": "assistant", "content": "<answer>done</answer>"})

            benchmark = SpreadsheetBench(root, {})
            benchmark.validate_task(task)
            result = InferenceAgent(Model(), benchmark=benchmark).run(task.public(), {}, directory=directory)
            self.assertEqual(benchmark.score("done", task, result), 1)
            workbook = openpyxl.load_workbook(result.artifacts["case-1/output.xlsx"], data_only=True)
            self.assertEqual(workbook.active["A1"].value, 5)
            workbook.close()
            ws = Workspace(root / "run")
            ws.initialize({"fixture": True})
            stored = ws.archive_artifacts(result.artifacts, directory, "raw/artifacts/fixture")
            self.assertTrue(ws.path(stored["case-1/output.xlsx"]["path"]).is_file())

    def test_timeout_and_read_only_container_data_boundary(self):
        with tempfile.TemporaryDirectory() as temporary:
            sandbox = DockerSandbox(Path(temporary), {"command_timeout": 2})
            sandbox.check()
            readonly = sandbox.run("touch /outside-task-file")
            self.assertNotEqual(readonly["exit_code"], 0)
            timed = sandbox.run("sleep 10")
            self.assertTrue(timed["timed_out"])
            self.assertEqual(timed["exit_code"], 124)


if __name__ == "__main__":
    unittest.main()

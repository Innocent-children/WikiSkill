from __future__ import annotations

import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from wikiskill.agents import Conversation, InferenceAgent, prompt
from wikiskill.benchmarks.alfworld import ALFWorld
from wikiskill.benchmarks.qa import GoogleSearch, LiveMath, OfficeQA, SealQA
from wikiskill.benchmarks.spreadsheet import SpreadsheetBench, compare_workbooks, answer_ranges
from wikiskill.data import Task, TaskInput
from wikiskill.demo import ScriptedModel
from wikiskill.model import ChatResponse
from wikiskill.storage import Workspace


class QueueModel:
    identity = {"adapter": "test"}

    def __init__(self, responses):
        self.responses = iter(responses)
        self.requests = []

    def complete(self, messages, tools):
        self.requests.append((messages, tools))
        return next(self.responses)


def answer(text):
    return ChatResponse({"role": "assistant", "content": text})


class QABenchmarkTests(unittest.TestCase):
    def test_original_prompt_hashes_and_single_step_math(self):
        from importlib.resources import files
        manifest = json.loads(files("wikiskill").joinpath("prompts/sources.json").read_text())
        self.assertEqual(len(manifest["prompts"]), 7)
        for filename, metadata in manifest["prompts"].items():
            content = files("wikiskill").joinpath("prompts/" + filename).read_bytes()
            self.assertEqual(hashlib.sha256(content).hexdigest(), metadata["sha256"])
        model = QueueModel([answer("<answer>B</answer>")])
        benchmark = LiveMath(Path("."), {})
        agent = InferenceAgent(model, benchmark=benchmark)
        task = Task("math", "Which is two?", "B", context={"choices": {"A": "1", "B": "2"}})
        result = agent.run(task.public(), {}, directory=Path("."))
        self.assertEqual(model.requests[0][0][0]["content"], prompt("live_math").replace("{skill_section}", ""))
        self.assertEqual(model.requests[0][1], [])
        self.assertNotIn("{skill_section}", model.requests[0][0][0]["content"])
        self.assertIn("B. 2", model.requests[0][0][1]["content"])
        self.assertEqual(benchmark.score("B", task, result), 1)

    def test_sealqa_search_file_and_separate_official_judge(self):
        config = {"base_url": "http://localhost/v1", "model": "judge", "api_key_env": ""}
        benchmark = SealQA(Path("."), {"judge_model": config})
        inference = QueueModel([ScriptedModel.call("web_search", {"query": "document title"}, 1),
            ScriptedModel.call("read_file", {"path": "search/results-1.json"}, 2), answer("<answer>42</answer>")])
        judge = QueueModel([answer("A")])
        benchmark.judge = judge
        task = Task("seal", "Find the value", "SECRET_GOLD_VALUE")
        agent = InferenceAgent(inference, benchmark=benchmark)
        with patch.object(benchmark.search_client, "search", return_value={"query": "document title", "results": [{"url": "https://example.org", "snippet": "42"}]}):
            result = agent.run(task.public(), {}, directory=Path("."))
        self.assertEqual(benchmark.score("42", task, result), 1)
        self.assertNotIn("SECRET_GOLD_VALUE", json.dumps(inference.requests))
        self.assertIn("SECRET_GOLD_VALUE", json.dumps(judge.requests))
        self.assertEqual(result.details["grading"]["verdict"], "A")
        benchmark.judge = QueueModel([answer("Maybe A")])
        with self.assertRaisesRegex(RuntimeError, "exactly A, B or C"):
            benchmark.score("42", task, result)

    def test_google_query_uses_environment_credentials_and_preserves_results(self):
        search = GoogleSearch({"api_key_env": "SEARCH_TEST_KEY", "search_engine_id_env": "SEARCH_TEST_CX"})
        with patch.dict(os.environ, {"SEARCH_TEST_KEY": "test-key", "SEARCH_TEST_CX": "test-cx"}), patch(
            "wikiskill.benchmarks.qa.request_json", return_value={"items": [{"title": "A", "link": "https://example.org", "snippet": "Found"}]}
        ) as request:
            result = search.search("a & b", 3)
        self.assertIn("q=a+%26+b", request.call_args.args[0])
        self.assertIn("cx=test-cx", request.call_args.args[0])
        self.assertEqual(result["results"][0]["snippet"], "Found")

    def test_officeqa_oracle_retrieval_and_upstream_grading(self):
        task = Task("office", "What was the amount?", "543 million", files={
            "treasury.txt": "Bulletin\nThe amount was 543 million.\nAnother page."},
            context={"oracle_pages": [{"path": "treasury.txt", "start_line": 1, "end_line": 2}]})
        benchmark = OfficeQA(Path("."), {})
        benchmark.validate_task(task)
        model = QueueModel([ScriptedModel.call("glob", {"pattern": "*.txt"}, 1),
            ScriptedModel.call("grep", {"pattern": "amount"}, 2),
            ScriptedModel.call("read", {"path": "treasury.txt", "start_line": 2, "end_line": 2}, 3),
            answer("<answer>543 million</answer>")])
        result = InferenceAgent(model, benchmark=benchmark).run(task.public(), {}, directory=Path("."))
        self.assertIn("Initial reference pages", model.requests[0][0][1]["content"])
        self.assertEqual(benchmark.score("543 million", task, result), 1)
        self.assertEqual(benchmark.score("987 million", task, result), 0)
        date_task = Task("date", "What date?", "March 1977")
        self.assertEqual(benchmark.score("April 1977", date_task, result), 0)
        self.assertEqual(benchmark.score("March 1977", date_task, result), 1)


class SpreadsheetGradingTests(unittest.TestCase):
    def setUp(self):
        import openpyxl
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        workbook = openpyxl.Workbook()
        workbook.active.title = "Data, annual"
        workbook.active["A1"] = 1.234
        workbook.active["B1"] = ""
        workbook.save(self.root / "reference.xlsx")
        workbook.close()

    def tearDown(self):
        self.temporary.cleanup()

    def test_answer_ranges_values_and_missing_output(self):
        import openpyxl
        workbook = openpyxl.load_workbook(self.root / "reference.xlsx")
        workbook.active["A1"] = "1.23"
        workbook.active["B1"] = None
        workbook.save(self.root / "output.xlsx")
        workbook.close()
        self.assertEqual(compare_workbooks(self.root / "reference.xlsx", self.root / "output.xlsx", "'Data, annual'!A1:B1")[0], 1)
        self.assertEqual(compare_workbooks(self.root / "reference.xlsx", self.root / "missing.xlsx", "A1")[0], 0)
        self.assertEqual(list(answer_ranges("'Data, annual'!A1:B1,C2", "Data, annual")),
                         [("Data, annual", "A1:B1"), ("Data, annual", "C2")])

    def test_uncached_reference_formula_is_a_data_error(self):
        import openpyxl
        workbook = openpyxl.load_workbook(self.root / "reference.xlsx")
        workbook.active["A1"] = "=1+2"
        workbook.save(self.root / "reference.xlsx")
        workbook.save(self.root / "output.xlsx")
        workbook.close()
        with self.assertRaisesRegex(ValueError, "cached result"):
            compare_workbooks(self.root / "reference.xlsx", self.root / "output.xlsx", "A1")

    def test_hard_and_soft_case_grading_and_private_references(self):
        import shutil
        shutil.copyfile(self.root / "reference.xlsx", self.root / "output.xlsx")
        task = Task("sheet", "Edit values", "workbook", context={"answer_position": "A1"},
                    evaluation={"reference_workbooks": ["reference.xlsx", "reference.xlsx"]})
        conversation = Conversation([], "", "answer", details={"cases": [
            {"case": 1, "output_workbook": str(self.root / "output.xlsx")},
            {"case": 2, "output_workbook": str(self.root / "missing.xlsx")}]})
        self.assertEqual(SpreadsheetBench(self.root, {}).score("", task, conversation), 0)
        self.assertEqual(SpreadsheetBench(self.root, {"score_mode": "soft"}).score("", task, conversation), 0.5)
        self.assertFalse(hasattr(task.public(), "evaluation"))


class ALFWorldTests(unittest.TestCase):
    def test_admissible_actions_environment_reward_and_close(self):
        class Environment:
            closed = False
            actions = []
            def reset(self):
                return ["A room."], {"admissible_commands": [["take book"]]}
            def step(self, actions):
                self.actions.extend(actions)
                return ["Done."], [1], [True], {"admissible_commands": [[]], "won": [True]}
            def close(self):
                self.closed = True

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "game.tw-pddl").write_text("{}")
            env = Environment()
            model = QueueModel([answer("<action>invalid</action>"), answer("<think>take it</think><action>take book</action>")])
            benchmark = ALFWorld(root, {"max_steps": 3})
            task = Task("alf", "Take the book.", "success", context={"game_file": "game.tw-pddl"})
            with patch("wikiskill.benchmarks.alfworld.open_environment", return_value=env):
                result = InferenceAgent(model, benchmark=benchmark).run(task.public(), {}, directory=root)
            self.assertEqual(env.actions, ["take book"])
            self.assertTrue(env.closed)
            self.assertEqual(benchmark.score("", task, result), 1)
            self.assertEqual(result.details["steps"][0]["feedback"], "invalid_action")
            self.assertNotIn("{current_observation}", json.dumps(model.requests))

    @unittest.skipUnless(os.environ.get("WIKISKILL_ALFWORLD_GAME"), "Set WIKISKILL_ALFWORLD_GAME for real simulator integration")
    def test_real_alfworld_walkthrough(self):
        path = Path(os.environ["WIKISKILL_ALFWORLD_GAME"]).resolve()
        game = json.loads(path.read_text())
        model = QueueModel([answer(f"<action>{action}</action>") for action in game["walkthrough"]])
        benchmark = ALFWorld(path.parent, {"max_steps": max(50, len(game["walkthrough"]))})
        task = Task("actual-game", "Complete the described household task.", "success", context={"game_file": path.name})
        result = InferenceAgent(model, benchmark=benchmark).run(task.public(), {}, directory=path.parent)
        self.assertTrue(result.details["won"])
        self.assertEqual(benchmark.score("", task, result), 1)


if __name__ == "__main__":
    unittest.main()

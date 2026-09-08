"""SpreadsheetBench workbook execution, recalculation and answer-range evaluation."""

from __future__ import annotations

import datetime as dt
import re
import shutil
from pathlib import Path

from ..agents import AgentError, Conversation, prompt
from ..data import Task, TaskInput, json_text
from ..environment import Tool, object_schema
from .base import Benchmark, asset_path
from .sandbox import DockerSandbox


def spreadsheet_library():
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install spreadsheet support with pip install '.[spreadsheet]'") from exc
    return openpyxl


def normalized_cell(value):
    if isinstance(value, dt.datetime):
        days = (value - dt.datetime(1899, 12, 30)).total_seconds() / 86400
        return round(days, 0)
    if isinstance(value, dt.time):
        return str(value)[:-3]
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, str):
        if value == "":
            return None
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


def answer_ranges(position: str, default_sheet: str):
    # Commas and exclamation marks inside quoted worksheet names remain part of the name.
    pieces = re.findall(r"(?:'(?:[^']|'')*'|[^,])+", position)
    if not pieces or not position.strip():
        raise ValueError("answer_position must identify cells or cell ranges")
    for piece in pieces:
        match = re.fullmatch(r"\s*(?:(?:'((?:[^']|'')*)'|([^!]+))!)?\s*'?([A-Za-z]+[1-9][0-9]*(?::[A-Za-z]+[1-9][0-9]*)?)'?\s*", piece)
        if not match:
            raise ValueError(f"Invalid answer range: {piece}")
        sheet = match[1].replace("''", "'") if match[1] is not None else match[2] or default_sheet
        yield sheet, match[3].upper()


def compare_workbooks(reference: Path, output: Path, position: str) -> tuple[float, dict]:
    library = spreadsheet_library()
    if not output.is_file() or output.is_symlink():
        return 0.0, {"reason": "output_workbook_missing"}
    try:
        actual = library.load_workbook(output, data_only=True)
    except Exception:
        return 0.0, {"reason": "invalid_output_workbook"}
    expected = library.load_workbook(reference, data_only=True)
    formulas = library.load_workbook(reference, data_only=False)
    checked = 0
    try:
        for sheet, region in answer_ranges(position, expected.sheetnames[0]):
            if sheet not in expected:
                raise ValueError(f"Reference workbook lacks worksheet {sheet}")
            if sheet not in actual:
                return 0.0, {"reason": "worksheet_missing", "sheet": sheet}
            from openpyxl.utils.cell import range_boundaries
            first_col, first_row, last_col, last_row = range_boundaries(region)
            if first_col > last_col or first_row > last_row:
                raise ValueError("Answer range has reversed boundaries")
            for row in range(first_row, last_row + 1):
                for column in range(first_col, last_col + 1):
                    cell = expected[sheet].cell(row, column)
                    formula = formulas[sheet].cell(row, column)
                    if formula.data_type == "f" and cell.value is None:
                        raise ValueError("Reference formula has no cached result; recalculate reference workbooks before evaluation")
                    left = normalized_cell(cell.value)
                    right = normalized_cell(actual[sheet].cell(row, column).value)
                    checked += 1
                    if type(left) is not type(right) or left != right:
                        return 0.0, {"reason": "cell_value_difference", "sheet": sheet, "cell": cell.coordinate,
                                     "expected": str(left), "actual": str(right), "checked_cells": checked}
        return 1.0, {"checked_cells": checked}
    finally:
        actual.close()
        expected.close()
        formulas.close()


class SpreadsheetBench(Benchmark):
    name = "spreadsheet"
    description = "SpreadsheetBench spreadsheet manipulation tasks through Python and bash"
    allowed_options = {"image", "command_timeout", "memory", "cpus", "output_characters", "score_mode"}

    def validate_task(self, task: Task):
        spreadsheet_library()
        cases = task.context.get("cases")
        references = task.evaluation.get("reference_workbooks")
        if not isinstance(cases, list) or not cases or not isinstance(references, list) or len(cases) != len(references):
            raise ValueError("Spreadsheet tasks need matching public cases and private reference_workbooks")
        if task.context.get("instruction_type") not in ("Cell-Level Manipulation", "Sheet-Level Manipulation"):
            raise ValueError("Invalid spreadsheet instruction_type")
        if not isinstance(task.context.get("answer_position"), str):
            raise ValueError("Spreadsheet task requires answer_position")
        for case in cases:
            if not isinstance(case, dict) or set(case) != {"input_workbook"}:
                raise ValueError("Each spreadsheet case supplies input_workbook")
            self.check_asset(case["input_workbook"])
        for reference in references:
            self.check_asset(reference)
        if self.options.get("score_mode", "hard") not in ("hard", "soft"):
            raise ValueError("Spreadsheet score_mode must be hard or soft")

    def run(self, agent, task: TaskInput, skills, wiki, directory):
        library = spreadsheet_library()
        messages, usage, artifacts, cases = [], [], {}, []
        for number, case in enumerate(task.context["cases"], 1):
            folder = directory / f"case-{number}"
            folder.mkdir(parents=True, exist_ok=False)
            source = asset_path(self.data_root, case["input_workbook"])
            shutil.copyfile(source, folder / "input.xlsx")
            workbook = library.load_workbook(source, data_only=True, read_only=True)
            try:
                preview = {sheet.title: [[str(cell) if cell is not None else None for cell in row]
                    for row in sheet.iter_rows(max_row=5, values_only=True)] for sheet in workbook}
            finally:
                workbook.close()
            sandbox = DockerSandbox(folder, self.options)
            sandbox.check()
            user = {"working_directory": "/workspace/task", "instruction": task.prompt,
                    "spreadsheet_path": "/workspace/task/input.xlsx", "spreadsheet_content": preview,
                    "instruction_type": task.context["instruction_type"],
                    "answer_position": task.context["answer_position"], "output_path": "/workspace/task/output.xlsx"}
            tool = Tool("bash", "Execute a shell command inside this task's spreadsheet working directory.",
                        object_schema({"command": {"type": "string"}}, ["command"]), sandbox.run)
            try:
                result = agent.run_with_tools(task, skills, [tool], template=prompt("spreadsheet"),
                                              user_content=json_text(user), wiki=wiki)
            except AgentError as exc:
                raise AgentError(str(exc), messages + exc.messages, usage + exc.usage) from exc
            messages.extend(result.messages)
            usage.extend(result.usage)
            output = folder / "output.xlsx"
            chosen = output
            recalc = None
            failure = None
            if output.is_symlink():
                failure = "symlink_output"
            elif output.is_file():
                try:
                    check = library.load_workbook(output, read_only=True)
                    check.close()
                except Exception:
                    failure = "invalid_output_workbook"
            else:
                failure = "output_workbook_missing"
            if failure is None:
                recalc = sandbox.run("libreoffice -env:UserInstallation=file:///tmp/wikiskill-profile --headless --convert-to xlsx --outdir recalculated output.xlsx")
                converted = folder / "recalculated/output.xlsx"
                if recalc["exit_code"] != 0 or not converted.is_file():
                    failure = "output_recalculation_failed"
                else:
                    chosen = converted
            if chosen.is_file() and not chosen.is_symlink():
                artifacts[f"case-{number}/output.xlsx"] = str(chosen)
            cases.append({"case": number, "output_workbook": str(chosen),
                          "termination": result.termination, "recalculation": recalc, "failure": failure})
        return Conversation(messages, "Spreadsheet execution completed.", "answer", usage, artifacts,
                            {"cases": cases, "score_mode": self.options.get("score_mode", "hard")})

    def score(self, prediction, task: Task, conversation):
        scores, diagnostics = [], []
        for case, reference in zip(conversation.details["cases"], task.evaluation["reference_workbooks"]):
            if case.get("failure"):
                score, detail = 0.0, {"reason": case["failure"]}
            else:
                score, detail = compare_workbooks(asset_path(self.data_root, reference),
                    Path(case["output_workbook"]), task.context["answer_position"])
            scores.append(score)
            diagnostics.append({"case": case["case"], "score": score, **detail})
        conversation.details.update(case_results=diagnostics, soft_score=sum(scores) / len(scores),
                                    hard_score=float(all(score == 1 for score in scores)))
        return conversation.details[self.options.get("score_mode", "hard") + "_score"]
